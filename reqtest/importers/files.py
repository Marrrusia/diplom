"""Чтение документов и табличный импорт с явным сопоставлением полей."""
import csv
import json
import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook

from reqtest.models.domain import Requirement, TestCase


@dataclass
class Table:
    """Заголовки и строки для предварительного просмотра."""
    headers: list[str]
    rows: list[list[str]]


def read_table(path: Path, sheet: str | None = None, encoding: str = "utf-8-sig") -> Table:
    """Прочитать CSV или выбранный лист XLSX без выполнения формул."""
    if path.suffix.lower() == ".csv":
        with path.open(encoding=encoding, newline="") as stream:
            sample = stream.read(8192)
            stream.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            rows = list(csv.reader(stream, dialect))
    elif path.suffix.lower() == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            page = workbook[sheet] if sheet else workbook.active
            rows = []
            for cells in page.iter_rows():
                if any(cell.data_type == "f" for cell in cells):
                    raise ValueError("Формулы не поддерживаются: замените их значениями")
                rows.append(["" if c.value is None else str(c.value) for c in cells])
        finally:
            workbook.close()
    else:
        raise ValueError("Ожидается CSV или XLSX")
    rows = [row for row in rows if any(value.strip() for value in row)]
    if not rows:
        raise ValueError("Файл не содержит данных")
    headers = [h.strip() for h in rows[0]]
    if not all(headers) or len(set(headers)) != len(headers):
        raise ValueError("Заголовки должны быть непустыми и уникальными")
    if any(len(row) > len(headers) for row in rows[1:]):
        raise ValueError("В строке больше значений, чем заголовков")
    return Table(headers, [row + [""] * (len(headers) - len(row)) for row in rows[1:]])


def _records(table: Table, mapping: dict[str, str], required: set[str]):
    if not required <= mapping.keys():
        raise ValueError("Не сопоставлены обязательные поля: " + ", ".join(sorted(required - mapping.keys())))
    if any(column not in table.headers for column in mapping.values()):
        raise ValueError("Сопоставленная колонка отсутствует")
    if len(set(mapping.values())) != len(mapping):
        raise ValueError("Одна колонка сопоставлена нескольким полям")
    for number, row in enumerate(table.rows, 2):
        yield number, {field: row[table.headers.index(column)].strip() for field, column in mapping.items()}


def import_requirements_table(table: Table, mapping: dict[str, str], source: str) -> list[Requirement]:
    """Проверить все строки; при ошибке не возвращать частичный результат."""
    items = []
    for number, record in _records(table, mapping, {"id", "text"}):
        try:
            extras = {h: table.rows[number - 2][i] for i, h in enumerate(table.headers) if h not in mapping.values()}
            items.append(Requirement(**{**record, "source": record.get("source") or source, "metadata": extras}))
        except ValueError as error:
            raise ValueError(f"Строка {number}: {error}") from error
    _unique(items)
    return items


def _lines(value: str) -> list[str]:
    return [line.strip() for line in value.splitlines() if line.strip()]


def import_tests_table(table: Table, mapping: dict[str, str]) -> list[TestCase]:
    """Загрузить существующие тесты; ввод — JSON-объект либо исходный текст."""
    items = []
    required = {"id", "requirement_id", "title", "steps", "expected_result"}
    for number, record in _records(table, mapping, required):
        try:
            record["steps"] = _lines(record["steps"])
            record["preconditions"] = _lines(record.get("preconditions", ""))
            value = record.get("input_data", "")
            record["input_data"] = json.loads(value) if value.startswith("{") else ({"текст": value} if value else {})
            items.append(TestCase(**record))
        except ValueError as error:
            raise ValueError(f"Строка {number}: {error}") from error
    _unique(items)
    return items


def _unique(items: list[Requirement] | list[TestCase]) -> None:
    ids = [item.id for item in items]
    if len(ids) != len(set(ids)):
        raise ValueError("Повторяющиеся идентификаторы в файле")
    if not items:
        raise ValueError("Нет записей для импорта")


def import_requirements_document(path: Path) -> list[Requirement]:
    """TXT/DOCX: блоки ID: текст; без ID назначается REQ-0001 и далее.

    Пустая строка отделяет блоки; строка с новым ID также начинает блок.
    DOCX-таблицы поддерживаются как отдельные строки ID | текст.
    """
    if path.suffix.lower() == ".txt":
        text = path.read_text(encoding="utf-8-sig")
    elif path.suffix.lower() == ".docx":
        doc = Document(path)
        chunks = []
        for node in doc.element.body:
            if node.tag == qn("w:p"):
                chunks.append(Paragraph(node, doc).text)
            elif node.tag == qn("w:tbl"):
                for row in node.iter(qn("w:tr")):
                    cells = ["\n".join(Paragraph(p, doc).text for p in c.findall(qn("w:p")))
                             for c in row.findall(qn("w:tc"))]
                    chunks.extend([": ".join(cells), ""])
        text = "\n".join(chunks)
    else:
        raise ValueError("Ожидается TXT или DOCX")
    pattern = re.compile(r"^([A-Za-zА-Яа-я][\w.-]*)\s*:\s*(.*)$")
    blocks: list[tuple[str | None, str]] = []
    current_id = None
    lines: list[str] = []
    for line in [*text.splitlines(), ""]:
        match = pattern.match(line.strip())
        if not line.strip() or match:
            if lines:
                blocks.append((current_id, "\n".join(lines)))
            current_id, lines = (match.group(1), [match.group(2)]) if match else (None, [])
        else:
            lines.append(line.strip())
    reserved = {identifier for identifier, _ in blocks if identifier}
    items = []
    counter = 1
    for identifier, content in blocks:
        if identifier is None:
            while f"REQ-{counter:04d}" in reserved:
                counter += 1
            identifier = f"REQ-{counter:04d}"
            reserved.add(identifier)
            counter += 1
        items.append(Requirement(id=identifier, text=content, source=str(path)))
    _unique(items)
    return items
