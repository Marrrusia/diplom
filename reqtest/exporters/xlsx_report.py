"""Отчёт из шести листов с сохранением доказательств и ручных решений."""
import json
import os
import tempfile
from enum import Enum
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reqtest.models.domain import Project, TestSource, ReviewStatus
from reqtest.services.coverage import active_elements, active_tests, matrix_status, coverage_summary


class ReportError(ValueError):
    """Данные нельзя экспортировать без потери содержания."""


def _value(value):
    if isinstance(value, Enum):
        return value.value
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False, default=lambda item: item.model_dump(mode="json"))
    return value


def _append(sheet, values) -> None:
    prepared = [_value(value) for value in values]
    if any(isinstance(value, str) and len(value) > 32767 for value in prepared):
        raise ReportError('Текст ячейки превышает лимит Excel (32767 символов). Данные не обрезаны; сохраните SQLite-проект')
    sheet.append(prepared)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            if len(cell.value) > 32767:
                raise ReportError('Текст ячейки превышает лимит Excel (32767 символов). Данные не обрезаны; сохраните SQLite-проект')
            cell.data_type = 's'  # Пользовательский текст никогда не становится формулой Excel.


def export_report(project: Project, path: Path) -> None:
    """Сформировать XLSX во временном файле и заменить целевой только при успехе."""
    project = Project.model_validate_json(project.model_dump_json())
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets = {name: workbook.create_sheet(name) for name in
              ('Requirements', 'Elements', 'TestCases', 'CoverageMatrix', 'Summary', 'GenerationLog')}
    for name, items, fields in (
        ('Requirements', project.requirements, ['id', 'text', 'type', 'source', 'metadata']),
        ('Elements', project.elements, ['id', 'requirement_id', 'type', 'description', 'source_fragment', 'review_status', 'source_refs', 'review_notes', 'condition_ids', 'negated_condition_ids', 'analysis_blocked']),
        ('TestCases', active_tests(project), ['id', 'requirement_id', 'title', 'preconditions', 'input_data', 'steps',
                                             'expected_result', 'source', 'review_status', 'element_ids']),
        ('GenerationLog', project.generation_log, ['id', 'requirement_id', 'model', 'mode', 'test_case_ids', 'notes'])):
        _append(sheets[name], fields)
        for item in items:
            _append(sheets[name], [getattr(item, field) for field in fields])
    tests = active_tests(project)
    elements = active_elements(project)
    if len(tests) + 3 > 16384:
        raise ReportError('Матрица превышает число столбцов Excel')
    matrix = sheets['CoverageMatrix']
    _append(matrix, ['element_id', 'requirement_id', 'description', *[t.id for t in tests]])
    for element in elements:
        _append(matrix, [element.id, element.requirement_id, element.description,
                        *[matrix_status(project, element.id, t.id).value for t in tests]])
    # На том же листе хранится таблица решений: дополнительные листы не нужны.
    _append(matrix, [])
    _append(matrix, ['CoverageDecisions'])
    fields = ['element_id', 'test_case_id', 'status', 'rationale', 'confidence', 'needs_manual_review',
              'review_status', 'manual_status', 'manual_rationale']
    _append(matrix, fields)
    test_ids = {t.id for t in tests}
    element_ids = {e.id for e in elements}
    for link in project.coverage:
        if link.test_case_id in test_ids and link.element_id in element_ids:
            _append(matrix, [getattr(link, field) for field in fields])
    summary = sheets['Summary']
    current = coverage_summary(project)
    _append(summary, ['scope', 'initial_percent', 'final_percent', 'elements', 'confirmed', 'unconfirmed_or_uncovered'])
    baseline = project.initial_coverage
    for scope, stats in [('Всего', current['total']), *current['requirements'].items()]:
        initial = (baseline.total if scope == 'Всего' else baseline.requirements.get(scope)) if baseline else None
        _append(summary, [scope, initial.percent if initial else 'не зафиксировано',
                         stats['percent'] if stats['percent'] is not None else 'не определено', stats['elements'], stats['covered'], stats['uncovered']])
    _append(summary, [])
    generated = [t for t in project.test_cases if t.source == TestSource.GENERATED]
    for label, value in [
        ('Проект', project.name),
        ('Исходное покрытие зафиксировано', baseline.captured_at if baseline else 'Нет: исторический процент не восстанавливается из текущих связей'),
        ('Сгенерированных тестов всего', len(generated)),
        ('Принятых сгенерированных тестов', sum(t.review_status == ReviewStatus.ACCEPTED for t in generated)),
        ('Ожидающих проверки сгенерированных тестов', sum(t.review_status == ReviewStatus.PENDING for t in generated)),
        ('Отклонённых сгенерированных тестов', sum(t.review_status == ReviewStatus.REJECTED for t in generated)),
        ('Отклонённых элементов (исключены)', current['rejected_elements']),
        ('Спорных элементов (исключены до проверки)', current['blocked_elements']),
        ('Требования с неподтверждённой полнотой выделения', current['analysis_incomplete']),
        ('Замечания анализа требований', project.analysis_notes),
        ('Оценка', 'Подтверждённое покрытие активных выделенных элементов; не результат выполнения тестов'),
        ('Правило TestCases', 'Неотклонённые существующие и только принятые сгенерированные тесты')]:
        _append(summary, [label, value])
    for sheet in sheets.values():
        if sheet.max_row > 1048576:
            raise ReportError('Превышен лимит строк Excel')
        sheet.freeze_panes = 'D2' if sheet.title == 'CoverageMatrix' else 'A2'
        if sheet.title != 'CoverageMatrix':
            sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='335C81')
        for row in sheet:
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical='top')
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 28
    path = path.resolve()
    descriptor, temporary = tempfile.mkstemp(prefix='.reqtest-report-', suffix='.xlsx', dir=path.parent)
    os.close(descriptor)
    temporary_path = Path(temporary)
    try:
        workbook.save(temporary_path)
        os.replace(temporary_path, path)
    finally:
        workbook.close()
        temporary_path.unlink(missing_ok=True)
