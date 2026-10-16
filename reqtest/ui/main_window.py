"""Русскоязычный интерфейс импорта, анализа, генерации и ручной проверки."""
import logging
from pathlib import Path
from PySide6.QtGui import QAction
from reqtest.ui.file_worker import FileWorker
from PySide6.QtWidgets import (QLabel, QMainWindow, QTabWidget, QVBoxLayout, QWidget,
    QHBoxLayout, QFormLayout, QComboBox, QLineEdit, QProgressBar, QTextEdit, QPushButton, QTableWidget, QTableWidgetItem, QFileDialog, QMessageBox, QDialog, QInputDialog)
from openpyxl import load_workbook
from reqtest.models.domain import Project, LLMSettings
from reqtest.ui.analysis_worker import AnalysisWorker
from reqtest.ui.coverage_worker import CoverageWorker
from reqtest.ui.generation_worker import GenerationWorker
from reqtest.services.generation import GenerationMode, review_generated_test
from reqtest.models.domain import TestSource, CoverageStatus
from reqtest.services.review import review_test, review_element
from reqtest.ui.edit_dialogs import TestEditorDialog, ElementEditorDialog
from reqtest.ui.presentation import configure_table, fit_table, filter_table, readable_record, enhance_page
from reqtest.services.coverage import active_elements, active_tests, matrix_status, coverage_summary, review_link
from reqtest.importers.files import read_table, import_requirements_document, import_requirements_table, import_tests_table
from reqtest.services.project_import import merge_import, update_requirement_text
from reqtest.ui.import_dialog import MappingDialog, FIELDS


class MainWindow(QMainWindow):
    """Главное окно с разделами будущего рабочего процесса."""

    def __init__(self) -> None:
        super().__init__()
        self.setWindowTitle("Интеллектуальный анализ покрытия требований")
        self.resize(1100, 720)
        self.project = Project(name="Новый проект")
        self.tables = {}
        self.worker = None
        self.import_buttons = []
        self.search_fields = {}
        self.record_details = {}
        self.operation_name = ""
        self.operation_error = False
        self.project_path = None
        self.saved_payload = self.project.model_dump_json()
        tabs = QTabWidget()
        for title in ("Требования", "Проверяемые элементы", "Тест-кейсы", "Матрица покрытия",
                      "Генерация тестов", "Отчёты", "Настройки LLM"):
            page = QWidget()
            layout = QVBoxLayout(page)
            if title in ("Требования", "Тест-кейсы"):
                kind = "requirements" if title == "Требования" else "test_cases"
                button = QPushButton("Импортировать файл")
                button.clicked.connect(lambda checked=False, k=kind: self.import_file(k))
                layout.addWidget(button)
                self.import_buttons.append(button)
                if kind == "requirements":
                    edit = QPushButton("Исправить текст выбранного требования")
                    edit.clicked.connect(self.edit_requirement)
                    layout.addWidget(edit)
                    self.import_buttons.append(edit)
                if kind == "test_cases":
                    for label, action in (("Редактировать выбранный тест", self.edit_selected_test),
                        ("Принять выбранный тест", lambda: self.review_selected_test(True)),
                        ("Отклонить выбранный тест", lambda: self.review_selected_test(False))):
                        edit_button = QPushButton(label)
                        edit_button.clicked.connect(action)
                        layout.addWidget(edit_button)
                        self.import_buttons.append(edit_button)
                table = QTableWidget(0, len(FIELDS[kind]) + (2 if kind == "test_cases" else 0))
                table.setHorizontalHeaderLabels(list(FIELDS[kind].values()) + (["Источник", "Проверка"] if kind == "test_cases" else []))
                table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
                self.tables[kind] = table
                self.add_search(layout, table, kind)
                layout.addWidget(table)
                details = QTextEdit()
                details.setReadOnly(True)
                details.setPlaceholderText("Выберите строку для просмотра полного текста")
                self.record_details[kind] = details
                layout.addWidget(details)
                table.currentCellChanged.connect(lambda row, col, oldrow, oldcol, k=kind: self.show_record(k, row))
                if kind == "test_cases":
                    table.cellDoubleClicked.connect(lambda row, col: self.edit_selected_test())
                tabs.addTab(page, title)
                continue
            if title == "Генерация тестов":
                self.generation_mode = QComboBox()
                self.generation_mode.addItems([mode.value for mode in GenerationMode])
                layout.addWidget(self.generation_mode)
                self.generate_button = QPushButton("Сгенерировать тесты")
                self.generate_button.clicked.connect(self.start_generation)
                layout.addWidget(self.generate_button)
                self.import_buttons.append(self.generate_button)
                cancel = QPushButton("Отменить текущую операцию")
                cancel.clicked.connect(self.cancel_analysis)
                layout.addWidget(cancel)
                self.generated_table = QTableWidget(0, 3)
                self.generated_table.setHorizontalHeaderLabels(["ID", "Название", "Проверка"])
                self.generated_table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
                self.generated_table.cellClicked.connect(self.show_generated_test)
                self.add_search(layout, self.generated_table, "generated")
                layout.addWidget(self.generated_table)
                self.generated_table.currentCellChanged.connect(lambda row, col, oldrow, oldcol: self.show_generated_test(row))
                self.generation_details = QTextEdit()
                self.generation_details.setReadOnly(True)
                layout.addWidget(self.generation_details)
                for label, accepted in (("Принять выбранный тест", True), ("Отклонить выбранный тест", False)):
                    button = QPushButton(label)
                    button.clicked.connect(lambda checked=False, a=accepted: self.review_generated(a))
                    layout.addWidget(button)
                    self.import_buttons.append(button)
                edit_generated = QPushButton("Редактировать выбранный тест")
                edit_generated.clicked.connect(lambda: self.edit_selected_test(True))
                layout.addWidget(edit_generated)
                self.import_buttons.append(edit_generated)
                self.generation_log = QTextEdit()
                self.generation_log.setReadOnly(True)
                layout.addWidget(self.generation_log)
                tabs.addTab(page, title)
                continue
            if title == "Матрица покрытия":
                self.coverage_button = QPushButton("Анализировать покрытие")
                self.coverage_button.clicked.connect(self.start_coverage)
                layout.addWidget(self.coverage_button)
                self.matrix = QTableWidget()
                self.matrix.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
                self.matrix.cellClicked.connect(self.show_link)
                self.matrix.currentCellChanged.connect(lambda row, col, oldrow, oldcol: self.show_link(row, col))
                layout.addWidget(self.matrix)
                self.coverage_details = QTextEdit()
                self.coverage_details.setReadOnly(True)
                layout.addWidget(self.coverage_details)
                for label, accepted in (("Подтвердить выбранное решение", True), ("Отклонить выбранную связь", False)):
                    button = QPushButton(label)
                    button.clicked.connect(lambda checked=False, a=accepted: self.review_selected_link(a))
                    self.import_buttons.append(button)
                    layout.addWidget(button)
                self.summary_label = QLabel("Покрытие ещё не рассчитано")
                self.summary_label.setWordWrap(True)
                layout.addWidget(self.summary_label)
                tabs.addTab(page, title)
                continue
            if title == "Отчёты":
                for label, action in (("Сохранить проект", self.save_current_project),
                                      ("Сохранить проект как…", lambda: self.save_current_project(True)),
                                      ("Открыть проект…", self.open_project_file),
                                      ("Экспортировать XLSX…", self.export_current_report)):
                    button = QPushButton(label)
                    button.clicked.connect(action)
                    layout.addWidget(button)
                    self.import_buttons.append(button)
                self.report_info = QTextEdit()
                self.report_info.setReadOnly(True)
                self.report_info.setPlainText("SQLite сохраняет все данные проекта и ручные решения.\n"
                    "XLSX содержит Requirements, Elements, TestCases, CoverageMatrix, Summary, GenerationLog.\n"
                    "API-ключ и настройки LLM в проект не записываются.\n"
                    "Исходный процент фиксируется перед первой генерацией; для старых проектов он может быть неизвестен.")
                layout.addWidget(self.report_info)
                tabs.addTab(page, title)
                continue
            if title == "Настройки LLM":
                form = QFormLayout()
                self.provider = QComboBox()
                self.provider.addItems(["ollama", "openai"])
                self.server = QLineEdit("http://localhost:11434")
                self.model_name = QLineEdit("qwen2.5:7b")
                self.api_key = QLineEdit()
                self.api_key.setEchoMode(QLineEdit.EchoMode.Password)
                for name, widget in (("Тип сервера", self.provider), ("Адрес сервера", self.server),
                                     ("Модель", self.model_name), ("API-ключ (необязательно)", self.api_key)):
                    form.addRow(name, widget)
                layout.addLayout(form)
                layout.addWidget(QLabel("Ollama: http://localhost:11434. OpenAI API: адрес с /v1.\n"
                                        "Ключ и настройки хранятся только в текущем сеансе."))
                layout.addStretch()
                tabs.addTab(page, title)
                continue
            if title == "Проверяемые элементы":
                self.analyze_button = QPushButton("Выделить элементы всех требований")
                self.analyze_button.clicked.connect(self.start_analysis)
                layout.addWidget(self.analyze_button)
                self.cancel_button = QPushButton("Отменить анализ")
                self.cancel_button.setEnabled(False)
                self.cancel_button.clicked.connect(self.cancel_analysis)
                layout.addWidget(self.cancel_button)
                self.analysis_progress = QProgressBar()
                layout.addWidget(self.analysis_progress)
                self.elements_table = QTableWidget(0, 6)
                self.elements_table.setHorizontalHeaderLabels(["ID", "Требование", "Категория", "Описание", "Цитата", "Проверка"])
                self.elements_table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
                self.add_search(layout, self.elements_table, "elements")
                layout.addWidget(self.elements_table)
                self.element_details = QTextEdit()
                self.element_details.setReadOnly(True)
                layout.addWidget(self.element_details)
                self.elements_table.currentCellChanged.connect(lambda row, col, oldrow, oldcol: self.show_element(row))
                self.elements_table.cellDoubleClicked.connect(lambda row, col: self.edit_selected_element())
                for label, action in (("Редактировать выбранный элемент", self.edit_selected_element),
                    ("Принять выбранный элемент", lambda: self.review_selected_element(True)),
                    ("Отклонить выбранный элемент", lambda: self.review_selected_element(False))):
                    edit_button = QPushButton(label)
                    edit_button.clicked.connect(action)
                    layout.addWidget(edit_button)
                    self.import_buttons.append(edit_button)
                self.analysis_notes = QTextEdit()
                self.analysis_notes.setReadOnly(True)
                layout.addWidget(self.analysis_notes)
                tabs.addTab(page, title)
                continue
            label = QLabel(f"{title}: раздел будет реализован на следующих этапах.")
            label.setWordWrap(True)
            layout.addWidget(label)
            layout.addStretch()
            tabs.addTab(page, title)
        self.setCentralWidget(tabs)
        for i in range(tabs.count()):
            enhance_page(tabs.widget(i).layout())
        tables = [*self.tables.values(), self.elements_table, self.generated_table]
        for table in tables:
            configure_table(table)
        configure_table(self.matrix, True)
        # Перемещаем существующий индикатор в строку состояния: он виден на всех вкладках.
        self.analysis_progress.setParent(self)
        self.analysis_progress.setRange(0, 100)
        self.analysis_progress.setValue(0)
        self.analysis_progress.setMaximumWidth(280)
        self.analysis_progress.setFormat("Готово")
        self.statusBar().addPermanentWidget(self.analysis_progress)
        self.global_cancel = QPushButton("Отменить операцию")
        self.global_cancel.setEnabled(False)
        self.global_cancel.clicked.connect(self.cancel_analysis)
        self.statusBar().addPermanentWidget(self.global_cancel)
        self.file_status = QLabel("Новый проект — сохраните в SQLite")
        self.statusBar().addPermanentWidget(self.file_status)
        self.setup_file_menu()
        self.provider.currentTextChanged.connect(self.provider_changed)
        self.statusBar().showMessage("Готово к работе: импорт, анализ и ручная проверка")

    def import_file(self, kind: str) -> None:
        """Выбрать файл и настройки; передать обработку импортёрам и сервису."""
        if kind not in ('requirements', 'test_cases'):
            QMessageBox.warning(self, 'Ошибка импорта', 'Неизвестный раздел импорта')
            return
        filters = "Требования (*.txt *.docx *.xlsx)" if kind == "requirements" else "Тест-кейсы (*.csv *.xlsx)"
        filename, _ = QFileDialog.getOpenFileName(self, "Импорт требований" if kind == "requirements" else "Импорт тест-кейсов", "", filters)
        if not filename:
            return
        try:
            path = Path(filename)
            allowed = {'.txt', '.docx', '.xlsx'} if kind == 'requirements' else {'.csv', '.xlsx'}
            if path.suffix.lower() not in allowed:
                section = 'требований' if kind == 'requirements' else 'тест-кейсов'
                raise ValueError(f'Импорт {section}: выбран файл «{path.name}». '
                    f'Фактическое расширение: {path.suffix or "отсутствует"}. '
                    f'Поддерживаются: {", ".join(sorted(allowed))}. Выберите исходный файл, а не архив или ярлык.')
            if kind == 'requirements' and path.suffix.lower() in (".txt", ".docx"):
                items = import_requirements_document(path)
            else:
                sheet = None
                encoding = "utf-8-sig"
                if path.suffix.lower() == ".xlsx":
                    workbook = load_workbook(path, read_only=True)
                    try:
                        names = workbook.sheetnames
                    finally:
                        workbook.close()
                    sheet, ok = QInputDialog.getItem(self, "Лист Excel", "Выберите лист", names, 0, False)
                    if not ok:
                        return
                else:
                    encoding, ok = QInputDialog.getItem(self, "Кодировка CSV", "Кодировка", ["utf-8-sig", "cp1251"], 0, False)
                    if not ok:
                        return
                table = read_table(path, sheet, encoding)
                dialog = MappingDialog(table, kind, self)
                if dialog.exec() != QDialog.DialogCode.Accepted:
                    return
                items = (import_requirements_table(table, dialog.mapping(), filename)
                         if kind == "requirements" else import_tests_table(table, dialog.mapping()))
            self.project = merge_import(self.project, items, kind)
            self.refresh_tables()
            self.statusBar().showMessage(f"Добавлено записей: {len(items)}. Сохраните изменения проекта")
        except Exception as error:
            logging.exception("Ошибка импорта")
            QMessageBox.warning(self, "Ошибка импорта", str(error))

    def refresh_tables(self) -> None:
        """Отобразить принятые записи проекта."""
        for kind, table in self.tables.items():
            items = getattr(self.project, kind)
            table.setRowCount(len(items))
            for i, item in enumerate(items):
                for j, field in enumerate([*FIELDS[kind], *(["source", "review_status"] if kind == "test_cases" else [])]):
                    value = getattr(item, field)
                    if hasattr(value, "value"):
                        value = value.value
                    if isinstance(value, list):
                        value = "\n".join(value)
                    elif isinstance(value, dict):
                        value = "\n".join(f"{key}: {entry}" for key, entry in value.items())
                    cell = QTableWidgetItem(str(value or ""))
                    cell.setToolTip(str(value or ""))
                    table.setItem(i, j, cell)
            fit_table(table)
            filter_table(table, self.search_fields[kind].text())
        self.render_matrix()
        self.render_generation()
        self.render_elements()
        for kind, table in self.tables.items():
            self.show_record(kind, table.currentRow())
        self.analysis_notes.setPlainText("Результаты модели требуют ручной проверки.\n" +
            "\n".join(f"{rid}: {note}" for rid, notes in self.project.analysis_notes.items() for note in notes))

    def edit_requirement(self) -> None:
        """Редактировать выбранное требование с проверкой снимка проекта."""
        row = self.tables["requirements"].currentRow()
        if row < 0:
            QMessageBox.information(self, "Выбор требования", "Выберите строку требования")
            return
        requirement = self.project.requirements[row]
        text, ok = QInputDialog.getMultiLineText(self, "Исправить текст", requirement.id, requirement.text)
        if not ok:
            return
        try:
            self.project = update_requirement_text(self.project, requirement.id, text)
            self.refresh_tables()
            self.render_elements()
            self.analysis_notes.clear()
        except ValueError as error:
            QMessageBox.warning(self, "Ошибка изменения", str(error))

    def start_analysis(self) -> None:
        """Собрать настройки сеанса и запустить фоновый анализ."""
        if self.worker is not None:
            return
        if not self.project.requirements:
            QMessageBox.information(self, "Анализ", "Сначала импортируйте требования")
            return
        if self.project.elements:
            answer = QMessageBox.question(self, "Повторный анализ",
                "Существующие элементы, связи тестов с элементами и покрытие будут сброшены. Продолжить?")
            if answer != QMessageBox.StandardButton.Yes:
                return
        settings = LLMSettings(provider=self.provider.currentText(), base_url=self.server.text(),
                               model=self.model_name.text(), api_key=self.api_key.text() or None)
        self.worker = AnalysisWorker(self.project, settings, self)
        self.coverage_button.setEnabled(False)
        self.worker.result.connect(self.accept_analysis)
        self.worker.error.connect(self.show_operation_error)
        self.worker.progress.connect(self.update_progress)
        self.worker.finished.connect(self.analysis_finished)
        self.analyze_button.setEnabled(False)
        self.cancel_button.setEnabled(True)
        for button in self.import_buttons:
            button.setEnabled(False)
        self.analysis_progress.setRange(0, len(self.project.requirements))
        self.analysis_progress.setValue(0)
        self.statusBar().showMessage("Анализ выполняется в фоне…")
        self.begin_operation()
        self.worker.start()

    def update_progress(self, value: int, total: int) -> None:
        """Обновить прогресс в GUI-потоке."""
        self.analysis_progress.setRange(0, total)
        self.analysis_progress.setValue(value)
        self.analysis_progress.setFormat(f"{self.operation_name}: %v/%m")

    def accept_analysis(self, result) -> None:
        """Принять только завершённый проверенный снимок."""
        self.project = result.project
        self.refresh_tables()
        self.analysis_notes.setPlainText("Результат LLM требует ручной проверки. Цитаты проверены автоматически; "
            "соответствие описаний смыслу требования не гарантируется.\n\n" +
            "\n".join(f"{identifier}: {note}" for identifier, notes in result.notes.items() for note in notes))

    def cancel_analysis(self) -> None:
        """Отменить между запросами; текущий HTTP-запрос ограничен таймаутом."""
        if isinstance(self.worker, FileWorker):
            self.statusBar().showMessage("Дождитесь завершения файловой операции")
            return
        if self.worker:
            self.worker.requestInterruption()
            self.cancel_button.setEnabled(False)
            self.global_cancel.setEnabled(False)
            self.statusBar().showMessage("Отмена после завершения текущего запроса (до 180 секунд)")

    def analysis_finished(self) -> None:
        """Освободить поток и разрешить следующую операцию."""
        self.worker.deleteLater()
        self.worker = None
        self.analyze_button.setEnabled(True)
        self.coverage_button.setEnabled(True)
        self.cancel_button.setEnabled(False)
        for button in self.import_buttons:
            button.setEnabled(True)
        self.global_cancel.setEnabled(False)
        for field in (self.provider, self.server, self.model_name, self.api_key):
            field.setEnabled(True)
        self.analysis_progress.setRange(0, 100)
        self.analysis_progress.setValue(0 if self.operation_error else 100)
        self.analysis_progress.setFormat("Ошибка / отмена" if self.operation_error else "Завершено")
        for action in self.file_actions:
            action.setEnabled(True)
        self.update_file_status()
        self.statusBar().showMessage("Операция не завершена" if self.operation_error else "Операция завершена")

    def closeEvent(self, event) -> None:
        """Не уничтожать работающий поток при закрытии окна."""
        if self.worker is not None:
            self.cancel_analysis()
            event.ignore()
            return
        if not self.confirm_discard():
            event.ignore()
            return
        event.accept()

    def render_elements(self) -> None:
        """Показать текущие предложения модели."""
        self.elements_table.setRowCount(len(self.project.elements))
        for i, element in enumerate(self.project.elements):
            values = (element.id, element.requirement_id, element.type.value,
                      element.description, element.source_fragment, element.review_status.value)
            for j, value in enumerate(values):
                item = QTableWidgetItem(value)
                item.setToolTip(value)
                self.elements_table.setItem(i, j, item)

        fit_table(self.elements_table)
        filter_table(self.elements_table, self.search_fields["elements"].text())
        self.show_element(self.elements_table.currentRow())

    def start_coverage(self) -> None:
        """Запустить атомарный анализ покрытия в фоне."""
        if self.worker is not None:
            return
        if not self.project.elements:
            QMessageBox.information(self, "Покрытие", "Сначала выделите элементы требований")
            return
        if self.project.coverage:
            if QMessageBox.question(self, "Повторный анализ", "Решения покрытия и их подтверждения будут заменены. Продолжить?") != QMessageBox.StandardButton.Yes:
                return
        settings = LLMSettings(provider=self.provider.currentText(), base_url=self.server.text(),
            model=self.model_name.text(), api_key=self.api_key.text() or None)
        self.worker = CoverageWorker(self.project, settings, self)
        self.worker.result.connect(self.accept_coverage)
        self.worker.error.connect(self.show_operation_error)
        self.worker.progress.connect(self.update_progress)
        self.worker.finished.connect(self.analysis_finished)
        self.coverage_button.setEnabled(False)
        self.analyze_button.setEnabled(False)
        self.cancel_button.setEnabled(True)
        for button in self.import_buttons:
            button.setEnabled(False)
        self.analysis_progress.setRange(0, 0)
        self.statusBar().showMessage("Анализ покрытия выполняется в фоне. Отмена — в разделе элементов")
        self.begin_operation()
        self.worker.start()

    def accept_coverage(self, project) -> None:
        """Показать проверенный результат без автоматического утверждения."""
        self.project = project
        self.render_matrix()

    def render_matrix(self) -> None:
        """Показать статусы и статистику только подтверждённого покрытия."""
        tests = active_tests(self.project)
        elements = active_elements(self.project)
        self.matrix.setRowCount(len(elements))
        self.matrix.setColumnCount(len(tests))
        self.matrix.setHorizontalHeaderLabels([t.id for t in tests])
        self.matrix.setVerticalHeaderLabels([e.id for e in elements])
        for i, element in enumerate(elements):
            for j, test in enumerate(tests):
                status = matrix_status(self.project, element.id, test.id)
                item = QTableWidgetItem(status.value)
                item.setToolTip(element.description + "\nТест: " + test.title)
                self.matrix.setItem(i, j, item)
        fit_table(self.matrix, True)
        stats = coverage_summary(self.project)
        lines = []
        for name, row in [("Всего", stats['total']), *stats['requirements'].items()]:
            percent = "не определено" if row['percent'] is None else f"{row['percent']}%"
            lines.append(f"{name}: подтверждено {row['covered']}/{row['elements']}; покрытие {percent}; неподтверждённых/непокрытых {row['uncovered']}")
        lines.append(f"Отклонённых элементов (исключены из расчёта): {stats['rejected_elements']}")
        self.summary_label.setText("\n".join(lines))
        self.show_link(self.matrix.currentRow(), self.matrix.currentColumn())
        if hasattr(self, "file_status"):
            self.update_file_status()

    def show_link(self, row: int, column: int) -> None:
        """Показать обоснование модели и техническую проверку выбранной пары."""
        if not 0 <= row < len(active_elements(self.project)) or not 0 <= column < len(active_tests(self.project)):
            self.coverage_details.clear()
            return
        element = active_elements(self.project)[row]
        test = active_tests(self.project)[column]
        link = next((l for l in self.project.coverage if l.element_id == element.id and l.test_case_id == test.id), None)
        text = f"{element.id} — {test.id}\n"
        if link:
            text += f"Предложение модели: {link.status.value}\nУверенность: {link.confidence}\nПроверка: {link.review_status.value}\n{link.rationale}"
            if link.manual_status:
                text += f"\nРучное решение: {link.manual_status.value}\n{link.manual_rationale}"
        else:
            text += "Нет решения: пара ещё не проанализирована или требования различаются"
        self.coverage_details.setPlainText(text)

    def review_selected_link(self, accepted: bool) -> None:
        """Передать ручное решение сервису и обновить процент."""
        row, column = self.matrix.currentRow(), self.matrix.currentColumn()
        if row < 0 or column < 0:
            QMessageBox.information(self, "Выбор", "Выберите ячейку матрицы")
            return
        manual_status = None
        manual_rationale = ""
        if accepted:
            status, ok = QInputDialog.getItem(self, "Ручная проверка", "Ваше решение", [s.value for s in CoverageStatus], 0, False)
            if not ok:
                return
            manual_rationale, ok = QInputDialog.getMultiLineText(self, "Обоснование", "Какие предусловия, действия и результаты подтверждают решение?")
            if not ok:
                return
            manual_status = CoverageStatus(status)
        try:
            self.project = review_link(self.project, active_elements(self.project)[row].id,
                active_tests(self.project)[column].id, accepted, manual_status, manual_rationale)
            self.render_matrix()
            self.show_link(row, column)
        except ValueError as error:
            QMessageBox.warning(self, "Проверка связи", str(error))

    def start_generation(self) -> None:
        """Начать фоновую генерацию в выбранном режиме."""
        if self.worker is not None:
            return
        settings = LLMSettings(provider=self.provider.currentText(), base_url=self.server.text(),
            model=self.model_name.text(), api_key=self.api_key.text() or None)
        mode = GenerationMode(self.generation_mode.currentText())
        self.worker = GenerationWorker(self.project, settings, mode, self)
        self.worker.result.connect(self.accept_generation)
        self.worker.error.connect(self.show_operation_error)
        self.worker.progress.connect(self.update_progress)
        self.worker.finished.connect(self.analysis_finished)
        self.analyze_button.setEnabled(False)
        self.coverage_button.setEnabled(False)
        self.cancel_button.setEnabled(True)
        for button in self.import_buttons:
            button.setEnabled(False)
        self.analysis_progress.setRange(0, 0)
        self.statusBar().showMessage("Генерация выполняется в фоне…")
        self.begin_operation()
        self.worker.start()

    def accept_generation(self, result) -> None:
        """Добавить предложения и показать итог, не подтверждая покрытие."""
        self.project = result.project
        self.refresh_tables()
        self.generation_details.setPlainText(result.message)

    def render_generation(self) -> None:
        """Показать предложения и журнал, включая причины отклонения."""
        tests = [t for t in self.project.test_cases if t.source == TestSource.GENERATED]
        self.generated_table.setRowCount(len(tests))
        for i, test in enumerate(tests):
            for j, value in enumerate((test.id, test.title, test.review_status.value)):
                self.generated_table.setItem(i, j, QTableWidgetItem(value))
        fit_table(self.generated_table)
        filter_table(self.generated_table, self.search_fields["generated"].text())
        self.generation_log.setPlainText("\n\n".join(
            f"{entry.requirement_id}: {entry.mode}; модель {entry.model}\n{entry.notes}"
            for entry in self.project.generation_log))

    def show_generated_test(self, row: int, column: int = 0) -> None:
        """Показать весь сценарий и его связи."""
        tests = [t for t in self.project.test_cases if t.source == TestSource.GENERATED]
        if 0 <= row < len(tests):
            self.generation_details.setPlainText(readable_record(tests[row], self.test_labels()))
        else:
            self.generation_details.clear()

    def review_generated(self, accepted: bool) -> None:
        """Принять/отклонить предложение без автоматического подтверждения покрытия."""
        row = self.generated_table.currentRow()
        tests = [t for t in self.project.test_cases if t.source == TestSource.GENERATED]
        if not 0 <= row < len(tests):
            QMessageBox.information(self, "Выбор", "Выберите сгенерированный тест")
            return
        self.project = review_test(self.project, tests[row].id, accepted)
        self.refresh_tables()
        self.show_generated_test(row)

    def edit_selected_test(self, generated: bool = False) -> None:
        """Открыть форму для существующего или сгенерированного теста."""
        if self.worker is not None:
            return
        tests = [t for t in self.project.test_cases if t.source == TestSource.GENERATED] if generated else self.project.test_cases
        table = self.generated_table if generated else self.tables['test_cases']
        row = table.currentRow()
        if not 0 <= row < len(tests):
            QMessageBox.information(self, 'Выбор', 'Выберите тест')
            return
        dialog = TestEditorDialog(self.project, tests[row], self)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            self.project = dialog.result_project
            self.refresh_tables()
            if generated:
                self.show_generated_test(row)

    def review_selected_test(self, accepted: bool) -> None:
        """Изменить статус выбранного теста, не подтверждая его покрытие."""
        if self.worker is not None:
            return
        row = self.tables['test_cases'].currentRow()
        if not 0 <= row < len(self.project.test_cases):
            QMessageBox.information(self, 'Выбор', 'Выберите тест')
            return
        self.project = review_test(self.project, self.project.test_cases[row].id, accepted)
        self.refresh_tables()

    def edit_selected_element(self) -> None:
        """Открыть проверяемый редактор элемента."""
        if self.worker is not None:
            return
        row = self.elements_table.currentRow()
        if not 0 <= row < len(self.project.elements):
            QMessageBox.information(self, 'Выбор', 'Выберите элемент')
            return
        dialog = ElementEditorDialog(self.project, self.project.elements[row], self)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            self.project = dialog.result_project
            self.refresh_tables()

    def review_selected_element(self, accepted: bool) -> None:
        """Изменить статус аспекта и сбросить зависимые решения."""
        if self.worker is not None:
            return
        row = self.elements_table.currentRow()
        if not 0 <= row < len(self.project.elements):
            QMessageBox.information(self, 'Выбор', 'Выберите элемент')
            return
        self.project = review_element(self.project, self.project.elements[row].id, accepted)
        self.refresh_tables()

    def add_search(self, layout, table: QTableWidget, kind: str) -> None:
        """Поиск в таблице без изменения моделей или порядка строк."""
        search = QLineEdit()
        search.setPlaceholderText("Поиск по ID и содержимому…")
        search.setClearButtonEnabled(True)
        search.textChanged.connect(lambda query: filter_table(table, query))
        self.search_fields[kind] = search
        layout.addWidget(search)

    @staticmethod
    def test_labels() -> dict[str, str]:
        """Русские подписи полного просмотра теста."""
        return {'id': 'Идентификатор', 'requirement_id': 'Требование', 'title': 'Название',
            'preconditions': 'Предусловия', 'input_data': 'Входные данные', 'steps': 'Действия',
            'expected_result': 'Ожидаемый результат', 'element_ids': 'Связанные элементы',
            'source': 'Источник', 'review_status': 'Проверка'}

    def show_record(self, kind: str, row: int) -> None:
        """Показать полную запись при выборе мышью или клавиатурой."""
        records = getattr(self.project, kind)
        if not 0 <= row < len(records):
            self.record_details[kind].clear()
            return
        labels = self.test_labels() if kind == 'test_cases' else {
            'id': 'Идентификатор', 'text': 'Текст требования', 'type': 'Тип', 'source': 'Источник', 'metadata': 'Дополнительные сведения'}
        self.record_details[kind].setPlainText(readable_record(records[row], labels))

    def show_element(self, row: int) -> None:
        """Показать описание и цитату полностью, вместе с исходным требованием."""
        if not 0 <= row < len(self.project.elements):
            self.element_details.clear()
            return
        element = self.project.elements[row]
        requirement = next(r for r in self.project.requirements if r.id == element.requirement_id)
        self.element_details.setPlainText(readable_record(element, {
            'id': 'Идентификатор', 'requirement_id': 'Требование', 'type': 'Категория',
            'description': 'Описание', 'source_fragment': 'Цитата', 'review_status': 'Проверка', 'review_notes': 'Замечания к структуре элемента', 'condition_ids': 'Применимо при условиях', 'negated_condition_ids': 'Применимо при ложности условий'}) +
            '\n\nИсходное требование:\n' + requirement.text +
            '\n\nСсылки на источники:\n' + ('\n'.join(f'{ref.source_id} [{ref.start}:{ref.end}]: {ref.text}' for ref in element.source_refs) or 'Старая версия или источник отредактирован вручную'))

    def begin_operation(self) -> None:
        """Единое состояние занятости и общая отмена на всех вкладках."""
        self.operation_error = False
        if isinstance(self.worker, FileWorker):
            self.operation_name = {'open': 'Открытие проекта', 'save': 'Сохранение проекта', 'export': 'Экспорт XLSX'}[self.worker.operation]
        elif isinstance(self.worker, GenerationWorker):
            self.operation_name = 'Генерация'
        elif isinstance(self.worker, CoverageWorker):
            self.operation_name = 'Покрытие'
        else:
            self.operation_name = 'Выделение элементов'
        self.global_cancel.setEnabled(not isinstance(self.worker, FileWorker))
        self.analysis_progress.setFormat(self.operation_name)
        self.statusBar().showMessage(self.operation_name + ' выполняется…')
        for field in (self.provider, self.server, self.model_name, self.api_key):
            field.setEnabled(False)
        for action in self.file_actions:
            action.setEnabled(False)

    def show_operation_error(self, message: str) -> None:
        """Отметить неуспешный исход, сохранив понятное объяснение модели/сервера."""
        self.operation_error = True
        QMessageBox.warning(self, 'Операция не выполнена', message)

    def provider_changed(self, provider: str) -> None:
        """Менять только стандартный адрес; сохранять пользовательский сервер."""
        if self.server.text() in ('http://localhost:11434', 'http://localhost:1234/v1'):
            self.server.setText('http://localhost:11434' if provider == 'ollama' else 'http://localhost:1234/v1')

    def setup_file_menu(self) -> None:
        """Файловые команды и стандартные сочетания клавиш."""
        menu = self.menuBar().addMenu("Файл")
        self.file_actions = []
        for title, shortcut, handler in (("Новый проект", "Ctrl+N", self.new_project),
            ("Открыть проект…", "Ctrl+O", self.open_project_file),
            ("Сохранить проект", "Ctrl+S", self.save_current_project),
            ("Сохранить проект как…", "Ctrl+Shift+S", lambda: self.save_current_project(True)),
            ("Экспорт XLSX…", "Ctrl+E", self.export_current_report)):
            action = QAction(title, self)
            action.setShortcut(shortcut)
            action.triggered.connect(handler)
            menu.addAction(action)
            self.file_actions.append(action)

    def update_file_status(self) -> None:
        """Показывать несохранённые изменения, сравнивая снимок с последней записью."""
        dirty = self.project.model_dump_json() != self.saved_payload
        name = self.project_path.name if self.project_path else "Новый проект"
        self.file_status.setText(name + (" — есть изменения" if dirty else ""))
        self.setWindowTitle("Интеллектуальный анализ покрытия требований — " + name + (" *" if dirty else ""))

    def confirm_discard(self) -> bool:
        """Предотвратить потерю изменений; сохранение выполняется отдельной командой."""
        if self.project.model_dump_json() == self.saved_payload:
            return True
        answer = QMessageBox.question(self, "Несохранённые изменения",
            "В проекте есть несохранённые изменения. Продолжить без сохранения?\n"
            "Выберите «Отмена», чтобы сначала сохранить проект (Ctrl+S).",
            QMessageBox.StandardButton.Discard | QMessageBox.StandardButton.Cancel,
            QMessageBox.StandardButton.Cancel)
        return answer == QMessageBox.StandardButton.Discard

    def new_project(self) -> None:
        """Начать отдельный проект после проверки несохранённых изменений."""
        if self.worker is not None or not self.confirm_discard():
            return
        name, ok = QInputDialog.getText(self, "Новый проект", "Название проекта")
        if not ok or not name.strip():
            return
        self.project = Project(name=name)
        self.project_path = None
        self.saved_payload = self.project.model_dump_json()
        self.clear_project_views()
        self.refresh_tables()

    def open_project_file(self) -> None:
        """Открыть файл без изменения текущего проекта до успешной проверки."""
        if self.worker is not None or not self.confirm_discard():
            return
        filename, _ = QFileDialog.getOpenFileName(self, "Открыть проект", "", "Проект SQLite (*.sqlite *.db)")
        if filename:
            self.start_file_operation('open', Path(filename))

    def save_current_project(self, save_as: bool = False) -> None:
        """Сохранить текущий снимок или выбрать новый файл."""
        if self.worker is not None:
            return
        path = None if save_as else self.project_path
        if path is None:
            filename, _ = QFileDialog.getSaveFileName(self, "Сохранить проект", "project.sqlite", "Проект SQLite (*.sqlite)")
            if not filename:
                return
            path = Path(filename)
            if not path.suffix:
                path = path.with_suffix('.sqlite')
        self.start_file_operation('save', path)

    def export_current_report(self) -> None:
        """Экспортировать отчёт, не заменяя SQLite-файл проекта."""
        if self.worker is not None:
            return
        filename, _ = QFileDialog.getSaveFileName(self, "Экспорт XLSX", "report.xlsx", "Отчёт Excel (*.xlsx)")
        if filename:
            path = Path(filename)
            if not path.suffix:
                path = path.with_suffix('.xlsx')
            if self.project_path and path.resolve() == self.project_path.resolve():
                QMessageBox.warning(self, 'Ошибка экспорта', 'Выберите отдельный файл отчёта')
                return
            self.start_file_operation('export', path)

    def start_file_operation(self, operation: str, path: Path) -> None:
        """Выполнить ввод/вывод в фоне с блокировкой изменения проекта."""
        self.worker = FileWorker(operation, path, self.project, self)
        self.worker.result.connect(self.accept_file_operation)
        self.worker.error.connect(self.show_operation_error)
        self.worker.finished.connect(self.analysis_finished)
        self.analyze_button.setEnabled(False)
        self.coverage_button.setEnabled(False)
        self.cancel_button.setEnabled(False)
        for button in self.import_buttons:
            button.setEnabled(False)
        self.analysis_progress.setRange(0, 0)
        self.begin_operation()
        self.worker.start()

    def accept_file_operation(self, result) -> None:
        """Отметить сохранение только после успешной записи на диск."""
        operation, path, project = result
        if operation in ('open', 'save'):
            self.project = project
            self.project_path = path
            self.saved_payload = project.model_dump_json()
            if operation == 'open':
                self.clear_project_views()
            self.refresh_tables()
        self.report_info.setPlainText(('Отчёт экспортирован: ' if operation == 'export' else 'Проект открыт: ' if operation == 'open' else 'Проект сохранён: ') + str(path))
        self.update_file_status()

    def clear_project_views(self) -> None:
        """Удалить текстовые результаты предыдущего проекта и фильтры."""
        for search in self.search_fields.values():
            search.clear()
        self.report_info.setPlainText("Сохраните проект в SQLite; отчёт экспортируется в отдельный XLSX-файл")
        for widget in [*self.record_details.values(), self.element_details, self.analysis_notes,
                       self.coverage_details, self.generation_details, self.generation_log]:
            widget.clear()
