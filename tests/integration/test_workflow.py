"""Полный рабочий цикл с реальными импортами/файлами и HTTP-контрактами LLM."""
import csv
import json
from pathlib import Path
import httpx
import pytest
from openpyxl import load_workbook
from reqtest.importers.files import import_requirements_document, read_table, import_tests_table
from reqtest.models.demo import demo_project
from reqtest.models.domain import Project, LLMSettings, CoverageStatus, ReviewStatus
from reqtest.services.project_import import merge_import
from reqtest.services.extraction import extract_project
from reqtest.services.coverage import analyze_coverage, review_link, coverage_summary
from reqtest.services.generation import generate_tests, GenerationMode
from reqtest.services.review import review_test, edit_test
from reqtest.storage.sqlite_project import save_project, load_project
from reqtest.exporters.xlsx_report import export_report
from reqtest.llm.client import LLMClient, LLMError


class ProtocolServer:
    """HTTP-подмена для обоих протоколов, с ответами на фактические запросы."""
    def __init__(self, provider: str) -> None:
        self.provider = provider
        self.requests = []

    def __call__(self, request: httpx.Request) -> httpx.Response:
        assert request.url.path == ('/api/chat' if self.provider == 'ollama' else '/v1/chat/completions')
        assert request.headers['Authorization'] == 'Bearer integration-key'
        body = json.loads(request.content)
        self.requests.append(body)
        messages = body['messages']
        prompt = messages[0]['content']
        payload = json.loads(messages[-1]['content'])
        if prompt.startswith('Ты аналитик'):
            result = {'elements': [{'type': e.type.value, 'description': e.description,
                'source_fragment': e.source_fragment} for e in demo_project().elements],
                'specification_gaps': ['Поведение при Instance != CRC не определено']}
        elif prompt.startswith('Оцени покрытие'):
            test, element = payload['test_case'], payload['element']
            # Импортированный частичный тест подтверждает только два условия.
            covered = test['id'] != 'EXISTING' or element['type'] == 'условие'
            result = {'status': 'покрыт' if covered else 'не покрыт', 'confidence': 0.95,
                'rationale': 'Сопоставлены предусловие, вызов и проверяемый результат',
                'precondition_evidence': test['preconditions'], 'step_evidence': test['steps'],
                'result_evidence': [test['expected_result']]}
        elif prompt.startswith('Сгенерируй тесты'):
            targets = payload['target_elements']
            result = {'tests': [{'title': 'Проверить включение тактирования CRC',
                'preconditions': ['crcHandle.Instance равно CRC'], 'input_data': {'crcHandle.Instance': 'CRC'},
                'steps': ['Вызвать HAL_CRC_MspInit',
                          'Зафиксировать вызов __HAL_RCC_CRC_CLK_ENABLE() компонентом без аргументов',
                          'Проверить включение тактирования CRC'],
                'expected_result': 'Вызван __HAL_RCC_CRC_CLK_ENABLE() без аргументов; тактирование CRC включено',
                'element_ids': [e['id'] for e in targets]}], 'specification_gaps': []}
        else:
            assert prompt.startswith('Проверь предложенный тест')
            result = {'consistent': True, 'sufficient_specification': True, 'duplicate': False,
                      'rationale': 'Проверка результата добавляет отсутствовавшие аспекты'}
        content = json.dumps(result, ensure_ascii=False)
        if self.provider == 'ollama':
            assert body['stream'] is False
            assert 'properties' in body['format']
            return httpx.Response(200, json={'message': {'content': content}})
        assert body['response_format'] == {'type': 'json_object'}
        return httpx.Response(200, json={'choices': [{'message': {'content': content}}]})


def confirm_proposals(project: Project) -> Project:
    """Явно принять положительные связи в имитации ручной проверки."""
    for link in list(project.coverage):
        if link.status == CoverageStatus.COVERED:
            project = review_link(project, link.element_id, link.test_case_id, True,
                CoverageStatus.COVERED, 'Проверены предусловия, действия и ожидаемый результат теста')
    return project


@pytest.mark.parametrize('provider', ['ollama', 'openai'])
@pytest.mark.parametrize('mode', [GenerationMode.FULL, GenerationMode.MISSING])
@pytest.mark.integration
def test_import_to_report_roundtrip(tmp_path: Path, provider: str, mode: GenerationMode):
    requirement = demo_project().requirements[0]
    source = tmp_path / 'Требования UTF-8.txt'
    source.write_text(f'{requirement.id}: {requirement.text}', encoding='utf-8-sig')
    project = merge_import(Project(name='Интеграционный CRC'), import_requirements_document(source), 'requirements')
    server = ProtocolServer(provider)
    client = LLMClient(LLMSettings(provider=provider, base_url='http://test-server' + ('/v1' if provider == 'openai' else ''),
                                 model='integration-model', api_key='integration-key'), httpx.MockTransport(server))
    project = extract_project(project, client).project
    assert len(project.elements) == 4
    assert project.analysis_notes[requirement.id]
    if mode == GenerationMode.MISSING:
        path = tmp_path / 'Тесты.csv'
        headers = ['id', 'requirement_id', 'title', 'preconditions', 'steps', 'expected_result']
        with path.open('w', encoding='utf-8-sig', newline='') as stream:
            writer = csv.writer(stream, delimiter=';')
            writer.writerows([headers, ['EXISTING', requirement.id, 'Проверка вызова', 'crcHandle.Instance равно CRC',
                                      'Вызвать HAL_CRC_MspInit', 'Вызов HAL_CRC_MspInit завершён']])
        project = merge_import(project, import_tests_table(read_table(path), {h: h for h in headers}), 'test_cases')
        project = confirm_proposals(analyze_coverage(project, client))
        assert coverage_summary(project)['total']['percent'] == 50
    initial = coverage_summary(project)['total']['percent']
    generated = generate_tests(project, client, mode)
    assert generated.generated_count == 1
    project = generated.project
    assert project.initial_coverage.total.percent == initial
    proposed = project.test_cases[-1]
    assert proposed.review_status == ReviewStatus.PENDING
    assert coverage_summary(project)['total']['percent'] == initial
    project = review_test(project, proposed.id, True)
    project = analyze_coverage(project, client)
    assert coverage_summary(project)['total']['percent'] == 0  # Повторный анализ сбросил подтверждения.
    project = confirm_proposals(project)
    assert coverage_summary(project)['total']['percent'] == 100
    saved = tmp_path / 'Сохранённый проект.sqlite'
    save_project(project, saved)
    restored = load_project(saved)
    assert restored == project
    assert restored.coverage[-1].manual_rationale
    assert 'integration-key' not in saved.read_bytes().decode('utf-8', errors='ignore')
    report = tmp_path / 'Итоговый отчёт.xlsx'
    export_report(restored, report)
    book = load_workbook(report)
    assert book['Summary']['B2'].value == initial
    assert book['Summary']['C2'].value == 100
    assert book['Summary']['F2'].value == 0
    assert len(book.sheetnames) == 6
    assert proposed.id in [row[0] for row in list(book['TestCases'].values)[1:]]
    book.close()
    requests_before = len(server.requests)
    assert generate_tests(restored, client, GenerationMode.MISSING).generated_count == 0
    assert len(server.requests) == requests_before
    changed = edit_test(restored, proposed.id, {'expected_result': proposed.expected_result + '; результат зафиксирован'})
    assert changed.test_cases[-1].review_status == ReviewStatus.PENDING
    assert not any(l.test_case_id == proposed.id for l in changed.coverage)
    assert changed.initial_coverage == restored.initial_coverage
