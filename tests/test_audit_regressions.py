"""Регрессии аудита: решения пользователя и блокировка фоновых операций."""
import pytest
from PySide6.QtWidgets import QApplication, QInputDialog, QFileDialog

from reqtest.models.domain import ReviewStatus
from reqtest.services.extraction import extract_project
from reqtest.services.project_import import update_requirement_text
from reqtest.services.review import edit_element, review_element, review_test
from test_review import reviewed_project
from test_llm import FakeClient, response


@pytest.mark.parametrize('operation', [
    lambda p: edit_element(p, p.elements[0].id, {'description': 'Вызов функции HAL_CRC_MspInit'}),
    lambda p: review_element(p, p.elements[0].id, False),
    lambda p: update_requirement_text(p, p.requirements[0].id, 'Новое требование'),
    lambda p: extract_project(p, FakeClient(response())).project,
])
def test_rejected_generated_test_stays_rejected(operation):
    original = review_test(reviewed_project(), 'T1', False)
    changed = operation(original)
    assert changed.test_cases[0].review_status == ReviewStatus.REJECTED
    assert original.test_cases[0].review_status == ReviewStatus.REJECTED
    assert not changed.coverage


@pytest.mark.parametrize('operation', ['edit_requirement', 'review_selected_link', 'review_generated', 'import_file'])
def test_busy_window_does_not_allow_project_changes(monkeypatch, operation):
    from reqtest.ui.main_window import MainWindow
    app = QApplication.instance() or QApplication([])
    window = MainWindow()
    window.project = reviewed_project()
    window.refresh_tables()
    window.tables['requirements'].setCurrentCell(0, 0)
    window.matrix.setCurrentCell(0, 0)
    window.generated_table.setCurrentCell(0, 0)
    before = window.project.model_dump_json()
    monkeypatch.setattr(QInputDialog, 'getMultiLineText',
                        lambda *args: pytest.fail('Редактирование во время фоновой операции'))
    monkeypatch.setattr(QFileDialog, 'getOpenFileName',
                        lambda *args: pytest.fail('Импорт во время фоновой операции'))
    window.worker = object()  # Состояние занятости без запуска реального потока.
    try:
        if operation == 'edit_requirement':
            window.edit_requirement()
        elif operation == 'import_file':
            window.import_file('requirements')
        else:
            getattr(window, operation)(False)
        assert window.project.model_dump_json() == before
    finally:
        window.worker = None
        window.close()


def test_docx_preserves_breaks_tabs_and_cell_paragraphs(tmp_path):
    from docx import Document
    from reqtest.importers.files import import_requirements_document
    path = tmp_path / 'requirements.docx'
    document = Document()
    document.add_paragraph('R1: Вызвать Foo\nЗатем проверить BAR\tбез аргументов')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'R2'
    table.cell(0, 1).text = 'Вызвать Init'
    table.cell(0, 1).add_paragraph('Проверить Ready')
    document.save(path)
    requirements = import_requirements_document(path)
    assert requirements[0].text == 'Вызвать Foo\nЗатем проверить BAR\tбез аргументов'
    assert requirements[1].text == 'Вызвать Init\nПроверить Ready'


@pytest.mark.parametrize('mode', ['первоначальный набор', 'дополнение покрытия'])
def test_rejected_elements_are_not_sent_as_covered(mode):
    import json
    from reqtest.models.demo import demo_project
    from reqtest.services.generation import generate_tests, GenerationMode
    from test_generation import Client
    project = review_element(demo_project(), 'E004', False)
    client = Client(tests=[])
    complete = client.complete

    def inspect(messages, schema):
        payload = json.loads(messages[-1]['content'])
        assert payload['covered_element_ids'] == []
        assert 'E004' not in [element['id'] for element in payload['target_elements']]
        return complete(messages, schema)

    client.complete = inspect
    generate_tests(project, client, GenerationMode(mode))
