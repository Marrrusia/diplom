"""SQLite и XLSX: реальные файлы, ручные решения, исходная статистика и ошибки."""
import sqlite3
import json
import time
import pytest
from openpyxl import load_workbook
from reqtest.models.domain import (Project, LLMSettings, ReviewStatus, CoverageStatus,
                                  TestCase as Case, TestSource as Source, CoverageLink, GenerationEntry)
from reqtest.models.demo import demo_project
from reqtest.services.generation import generate_tests, GenerationMode
from reqtest.services.coverage import review_link, coverage_summary
from reqtest.storage.sqlite_project import load_project, save_project, ProjectFileError
from reqtest.exporters.xlsx_report import export_report, ReportError


def rich_project():
    project = demo_project()
    data = project.model_dump()
    for identifier, source, status in [('T1', Source.EXISTING, ReviewStatus.PENDING),
        ('G1', Source.GENERATED, ReviewStatus.ACCEPTED), ('G2', Source.GENERATED, ReviewStatus.REJECTED),
        ('G3', Source.GENERATED, ReviewStatus.PENDING)]:
        data['test_cases'].append(Case(id=identifier, requirement_id=project.requirements[0].id,
            title='Русский тест ' + identifier, preconditions=['crcHandle.Instance == CRC'],
            input_data={'crcHandle.Instance': 'CRC'}, steps=['Вызвать HAL_CRC_MspInit'],
            expected_result='Тактирование CRC включено', source=source, review_status=status,
            element_ids=[e.id for e in project.elements]).model_dump())
    data['coverage'] = [CoverageLink(element_id=e.id, test_case_id='G1', status=CoverageStatus.REVIEW,
        rationale='Обоснование модели', confidence=0.7, review_status=ReviewStatus.ACCEPTED,
        manual_status=CoverageStatus.COVERED, manual_rationale='Действия и результат проверены пользователем').model_dump()
        for e in project.elements]
    data['generation_log'] = [GenerationEntry(id='L1', requirement_id=project.requirements[0].id,
        model='test', mode='full', test_case_ids=['G1', 'G2', 'G3'], notes='Историческая генерация').model_dump()]
    data['analysis_notes'] = {project.requirements[0].id: ['Поведение при Instance != CRC не определено']}
    data['initial_coverage'] = {'captured_at': '2026-10-14T00:00:00+00:00',
        'total': {'elements': 4, 'covered': 0}, 'requirements': {project.requirements[0].id: {'elements': 4, 'covered': 0}}}
    return Project.model_validate(data)


def test_sqlite_roundtrip_preserves_everything(tmp_path):
    path = tmp_path / 'Проект с пробелами.sqlite'
    original = rich_project()
    save_project(original, path)
    restored = load_project(path)
    assert restored == original
    assert coverage_summary(restored)['total']['percent'] == 100
    assert restored.coverage[0].manual_rationale
    save_project(restored, path)
    assert load_project(path) == original
    connection = sqlite3.connect(path)
    payload = connection.execute('SELECT payload FROM project_snapshot').fetchone()[0]
    connection.close()
    assert 'api_key' not in payload
    assert 'base_url' not in payload


def test_wrong_files_never_overwritten(tmp_path):
    path = tmp_path / 'other.sqlite'
    path.write_bytes(b'not a project')
    before = path.read_bytes()
    with pytest.raises(ProjectFileError):
        save_project(rich_project(), path)
    assert path.read_bytes() == before
    other = tmp_path / 'other.db'
    connection = sqlite3.connect(other)
    connection.execute('CREATE TABLE user_data (value TEXT)')
    connection.commit()
    connection.close()
    with pytest.raises(ProjectFileError, match='не является'):
        load_project(other)


def test_unknown_version_and_corrupted_payload_rejected(tmp_path):
    path = tmp_path / 'project.sqlite'
    save_project(rich_project(), path)
    connection = sqlite3.connect(path)
    connection.execute('PRAGMA user_version=999')
    connection.commit()
    connection.close()
    with pytest.raises(ProjectFileError, match='Версия'):
        load_project(path)
    connection = sqlite3.connect(path)
    connection.execute('PRAGMA user_version=1')
    connection.execute("UPDATE project_snapshot SET payload='{}'")
    connection.commit()
    connection.close()
    with pytest.raises(ProjectFileError):
        load_project(path)


def test_failed_atomic_save_preserves_previous_file(tmp_path, monkeypatch):
    from reqtest.storage import sqlite_project
    path = tmp_path / 'p.sqlite'
    original = rich_project()
    save_project(original, path)
    def fail(*args):
        raise PermissionError('Файл заблокирован')
    monkeypatch.setattr(sqlite_project.os, 'replace', fail)
    with pytest.raises(PermissionError):
        save_project(original.model_copy(update={'name': 'Новое имя'}), path)
    assert load_project(path) == original
    assert not list(tmp_path.glob('.reqtest-*'))


def test_export_six_sheets_and_approved_set(tmp_path):
    path = tmp_path / 'Отчёт.xlsx'
    export_report(rich_project(), path)
    book = load_workbook(path)
    assert book.sheetnames == ['Requirements', 'Elements', 'TestCases', 'CoverageMatrix', 'Summary', 'GenerationLog']
    assert [row[0].value for row in list(book['TestCases'].rows)[1:]] == ['T1', 'G1']
    assert book['CoverageMatrix']['E2'].value == 'покрыт'
    rows = list(book['Summary'].values)
    assert rows[1][1:6] == (0, 100, 4, 4, 0)
    assert any('Действия и результат проверены пользователем' in str(row) for row in book['CoverageMatrix'].values)
    assert 'Историческая генерация' in str(list(book['GenerationLog'].values))
    book.close()


def test_unknown_baseline_and_empty_denominator_not_fake_zero(tmp_path):
    path = tmp_path / 'empty.xlsx'
    export_report(Project(name='Пустой'), path)
    book = load_workbook(path)
    assert book['Summary']['B2'].value == 'не зафиксировано'
    assert book['Summary']['C2'].value == 'не определено'
    book.close()


def test_user_text_never_exported_as_formula(tmp_path):
    project = rich_project()
    project.requirements[0].metadata = {'formula': '=HYPERLINK("https://example.com")'}
    project.test_cases[0].title = '=1+1'
    path = tmp_path / 'safe.xlsx'
    export_report(project, path)
    book = load_workbook(path)
    cell = book['TestCases']['C2']
    assert cell.value == '=1+1' and cell.data_type == 's'
    book.close()


def test_long_cell_not_silently_truncated(tmp_path):
    project = rich_project()
    project.test_cases[0].expected_result = 'А' * 32768
    path = tmp_path / 'report.xlsx'
    path.write_bytes(b'previous')
    with pytest.raises(ReportError, match='32767'):
        export_report(project, path)
    assert path.read_bytes() == b'previous'


def test_baseline_captured_before_first_generation():
    class Client:
        settings = LLMSettings(model='test')
        def complete(self, messages, schema):
            return json.dumps({'tests': [], 'specification_gaps': ['Недостаточно информации']})
    generated = generate_tests(demo_project(), Client(), GenerationMode.FULL)
    assert generated.project.initial_coverage.total.percent == 0
    assert generated.project.initial_coverage.total.elements == 4
    assert generated.project.initial_coverage.captured_at


def test_gui_save_open_export_roundtrip(tmp_path, monkeypatch):
    from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox
    from reqtest.ui.main_window import MainWindow
    app = QApplication.instance() or QApplication([])
    window = MainWindow()
    window.project = rich_project()
    window.refresh_tables()
    path = tmp_path / 'gui.sqlite'
    monkeypatch.setattr(QFileDialog, 'getSaveFileName', lambda *args: (str(path), ''))
    def wait():
        deadline = time.monotonic() + 5
        while window.worker is not None and time.monotonic() < deadline:
            app.processEvents()
            time.sleep(0.005)
        assert window.worker is None
    window.save_current_project()
    wait()
    assert window.project_path == path
    assert window.project.model_dump_json() == window.saved_payload
    before = window.project.model_dump_json()
    window.project = Project(name='Другой')
    monkeypatch.setattr(QFileDialog, 'getOpenFileName', lambda *args: (str(path), ''))
    window.open_project_file()
    wait()
    assert window.project.model_dump_json() == before
    assert window.matrix.item(0, 1).text() == 'покрыт'
    output = tmp_path / 'gui.xlsx'
    monkeypatch.setattr(QFileDialog, 'getSaveFileName', lambda *args: (str(output), ''))
    window.export_current_report()
    wait()
    assert output.exists()
    window.close()


def test_unsaved_cancel_prevents_close(monkeypatch):
    from PySide6.QtWidgets import QApplication, QMessageBox
    from reqtest.ui.main_window import MainWindow
    app = QApplication.instance() or QApplication([])
    window = MainWindow()
    window.project = demo_project()
    window.show()
    monkeypatch.setattr(QMessageBox, 'question', lambda *args: QMessageBox.StandardButton.Cancel)
    assert not window.close()
    assert window.isVisible()
    monkeypatch.setattr(QMessageBox, 'question', lambda *args: QMessageBox.StandardButton.Discard)
    assert window.close()


def test_open_failure_keeps_current_project(tmp_path, monkeypatch):
    from PySide6.QtWidgets import QApplication, QMessageBox
    from reqtest.ui.main_window import MainWindow
    app = QApplication.instance() or QApplication([])
    window = MainWindow()
    window.project = rich_project()
    before = window.project.model_dump_json()
    invalid = tmp_path / 'broken.sqlite'
    invalid.write_bytes(b'broken')
    messages = []
    monkeypatch.setattr(QMessageBox, 'warning', lambda *args: messages.append(args[2]))
    window.start_file_operation('open', invalid)
    deadline = time.monotonic() + 5
    while window.worker is not None and time.monotonic() < deadline:
        app.processEvents()
        time.sleep(0.005)
    assert window.worker is None
    assert messages and window.operation_error
    assert window.project.model_dump_json() == before
    assert window.project_path is None
    window.close()
