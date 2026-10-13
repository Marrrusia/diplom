"""Реальные файлы всех форматов и атомарность импорта."""
import pytest
from docx import Document
from openpyxl import Workbook
from reqtest.importers.files import (Table, read_table, import_requirements_document,
    import_requirements_table, import_tests_table)
from reqtest.models.domain import Project
from reqtest.services.project_import import merge_import


def test_txt_multiline_and_ids(tmp_path):
    path = tmp_path / 'r.txt'
    path.write_text('R1: Условие\nпродолжение\n\nБез идентификатора', encoding='utf-8-sig')
    items = import_requirements_document(path)
    assert [r.id for r in items] == ['R1', 'REQ-0001']
    assert items[0].text == 'Условие\nпродолжение'


def test_docx_paragraph_and_table(tmp_path):
    path = tmp_path / 'r.docx'
    doc = Document()
    doc.add_paragraph('R1: Первый текст')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'R2'
    table.cell(0, 1).text = 'Второй текст'
    doc.save(path)
    assert [r.id for r in import_requirements_document(path)] == ['R1', 'R2']


def test_xlsx_mapping_metadata_sheet(tmp_path):
    path = tmp_path / 'r.xlsx'
    book = Workbook()
    page = book.create_sheet('Требования')
    page.append(['Номер', 'Описание', 'Примечание'])
    page.append(['R1', 'Русский текст', 'Дополнительно'])
    book.save(path)
    records = import_requirements_table(read_table(path, 'Требования'), {'id': 'Номер', 'text': 'Описание'}, str(path))
    assert records[0].metadata == {'Примечание': 'Дополнительно'}


@pytest.mark.parametrize('extension', ['csv', 'xlsx'])
def test_tests_import_and_merge(tmp_path, extension):
    path = tmp_path / ('tests.' + extension)
    headers = ['id', 'requirement_id', 'title', 'steps', 'expected_result']
    row = ['T1', 'R1', 'Тест', 'Шаг 1\nШаг 2', 'Результат']
    if extension == 'csv':
        import csv
        with path.open('w', encoding='utf-8-sig', newline='') as stream:
            writer = csv.writer(stream, delimiter=';')
            writer.writerows([headers, row])
    else:
        book = Workbook()
        book.active.append(headers)
        book.active.append(row)
        book.save(path)
    tests = import_tests_table(read_table(path), {h: h for h in headers})
    assert tests[0].steps == ['Шаг 1', 'Шаг 2']
    project = Project(name='Проект')
    with pytest.raises(ValueError):
        merge_import(project, tests, 'test_cases')
    assert not project.test_cases


def test_duplicate_and_missing_mapping():
    table = Table(['id', 'text'], [['R1', 'A'], ['R1', 'B']])
    with pytest.raises(ValueError, match='Повторяющиеся'):
        import_requirements_table(table, {'id': 'id', 'text': 'text'}, 'file')
    with pytest.raises(ValueError, match='обязательные'):
        import_requirements_table(table, {'id': 'id'}, 'file')


def test_formula_rejected(tmp_path):
    path = tmp_path / 'r.xlsx'
    book = Workbook()
    book.active.append(['id', 'text'])
    book.active.append(['R1', '=1+1'])
    book.save(path)
    with pytest.raises(ValueError, match='Формулы'):
        read_table(path)


def test_merge_duplicate_preserves_project():
    from reqtest.models.domain import Requirement
    r = Requirement(id='R1', text='Текст')
    project = Project(name='Проект', requirements=[r])
    with pytest.raises(ValueError):
        merge_import(project, [r], 'requirements')
    assert len(project.requirements) == 1


def test_edit_is_atomic():
    from reqtest.models.demo import demo_project
    from reqtest.services.project_import import update_requirement_text
    project = demo_project()
    with pytest.raises(ValueError):
        update_requirement_text(project, project.requirements[0].id, '')
    assert project.elements[0].source_fragment in project.requirements[0].text


def test_example_imports():
    from pathlib import Path
    root = Path(__file__).resolve().parents[1] / 'examples'
    requirements = import_requirements_document(root / 'requirements.txt')
    tests = import_tests_table(read_table(root / 'test_cases.csv'), {h: h for h in ['id','requirement_id','title','preconditions','input_data','steps','expected_result']})
    project = merge_import(Project(name='Пример', requirements=requirements), tests, 'test_cases')
    assert len(project.test_cases) == 1
